"""
YouTube Auto-SEO System — Multi-Channel

Updates:
- OAuth mein youtube.force-ssl scope included.
- Analytics mein views aur averageViewDuration fetch hoti hain.
- Unavailable impressions aur CTR "N/A" rehte hain.
- CTR unavailable ho to keyword learning skip hoti hai.

Har channel channels_config.json se load hota hai.
Prefixed secrets workflow ke env section mein bhi map hone chahiye.
"""

import os
import re
import json
import traceback
import requests
from datetime import datetime, timezone

from pytrends.request import TrendReq
from googleapiclient.discovery import build
from google.oauth2.credentials import Credentials
from openpyxl import Workbook, load_workbook

import growth_features as gf
import upload_features as uf
import quality_features as qf


# ---------- Shared Config ----------
GROQ_API_KEY = os.environ.get("GROQ_API_KEY")
WHATSAPP_PHONE = os.environ.get("WHATSAPP_PHONE")
WHATSAPP_APIKEY = os.environ.get("WHATSAPP_APIKEY")

CHANNELS_CONFIG_FILE = "channels_config.json"

LOW_CTR_THRESHOLD = 2.0
LOW_IMPRESSIONS_THRESHOLD = 500
LOW_RETENTION_PERCENT = 30
DIAGNOSIS_MIN_AGE_DAYS = 7


# ---------- Notifications ----------
def send_whatsapp_message(text):
    if not (WHATSAPP_PHONE and WHATSAPP_APIKEY):
        print(
            "WhatsApp configure nahi hai, "
            "skip kar rahe hain notification."
        )
        return

    try:
        requests.get(
            "https://api.callmebot.com/whatsapp.php",
            params={
                "phone": WHATSAPP_PHONE,
                "text": text[:2000],
                "apikey": WHATSAPP_APIKEY,
            },
            timeout=15,
        )
    except Exception as e:
        print(f"WhatsApp message bhejne mein masla hua: {e}")


# ---------- Channel Config Loading ----------
def load_channels():
    if not os.path.exists(CHANNELS_CONFIG_FILE):
        raise SystemExit(
            f"{CHANNELS_CONFIG_FILE} nahi mili "
            "— repo root mein honi chahiye."
        )

    with open(CHANNELS_CONFIG_FILE, "r", encoding="utf-8") as f:
        data = json.load(f)

    return data.get("channels", [])


def build_channel_context(channel_cfg):
    prefix = channel_cfg.get("secret_prefix", "")
    name = channel_cfg.get("name", prefix or "channel")

    ctx = {
        "name": name,
        "niche_keywords": channel_cfg.get(
            "niche_keywords",
            "your niche keywords here",
        ),
        "target_audience_geos": channel_cfg.get(
            "target_audience_geos",
            "US,GB",
        ).split(","),
        "target_audience_label": channel_cfg.get(
            "target_audience_label",
            "US and UK",
        ),
        "publish_hour_utc": int(
            channel_cfg.get("publish_hour_utc", 21)
        ),
        "shorts_upload_days": channel_cfg.get(
            "shorts_upload_days",
            "0,1,2,3,4",
        ),
        "long_upload_days": channel_cfg.get(
            "long_upload_days",
            "5",
        ),
        "enable_caption_improvement": bool(
            channel_cfg.get("enable_caption_improvement", True)
        ),
        "enable_auto_comment": bool(
            channel_cfg.get("enable_auto_comment", False)
        ),
        "enable_metadata_refresh": bool(
            channel_cfg.get("enable_metadata_refresh", False)
        ),
        "video_language": channel_cfg.get("video_language", "en"),
        "category_id": str(channel_cfg.get("category_id", "27")),
        "made_for_kids": bool(channel_cfg.get("made_for_kids", False)),
        "contains_synthetic_media": bool(
            channel_cfg.get("contains_synthetic_media", True)
        ),
        "quota_safety_limit": int(
            os.environ.get(
                f"{prefix}QUOTA_SAFETY_LIMIT",
                os.environ.get("QUOTA_SAFETY_LIMIT", "9000"),
            )
        ),
        "youtube_api_key": os.environ.get(
            f"{prefix}YOUTUBE_API_KEY"
        ),
        "youtube_channel_id": os.environ.get(
            f"{prefix}YOUTUBE_CHANNEL_ID"
        ),
        "oauth_client_id": os.environ.get(
            f"{prefix}YOUTUBE_OAUTH_CLIENT_ID"
        ),
        "oauth_client_secret": os.environ.get(
            f"{prefix}YOUTUBE_OAUTH_CLIENT_SECRET"
        ),
        "oauth_refresh_token": os.environ.get(
            f"{prefix}YOUTUBE_OAUTH_REFRESH_TOKEN"
        ),
        "drive_shorts_toupload": os.environ.get(
            f"{prefix}DRIVE_SHORTS_TOUPLOAD_FOLDER_ID"
        ),
        "drive_shorts_uploaded": os.environ.get(
            f"{prefix}DRIVE_SHORTS_UPLOADED_FOLDER_ID"
        ),
        "drive_long_toupload": os.environ.get(
            f"{prefix}DRIVE_LONG_TOUPLOAD_FOLDER_ID"
        ),
        "drive_long_uploaded": os.environ.get(
            f"{prefix}DRIVE_LONG_UPLOADED_FOLDER_ID"
        ),
        "tracking_file": f"tracking_sheet_{name}.xlsx",
        "winning_keywords_file": f"winning_keywords_{name}.json",
    }

    return ctx


# ---------- OAuth ----------
def get_oauth_credentials(ctx):
    return Credentials(
        token=None,
        refresh_token=ctx["oauth_refresh_token"],
        client_id=ctx["oauth_client_id"],
        client_secret=ctx["oauth_client_secret"],
        token_uri="https://oauth2.googleapis.com/token",
        scopes=[
            "https://www.googleapis.com/auth/youtube",
            "https://www.googleapis.com/auth/youtube.force-ssl",
            "https://www.googleapis.com/auth/yt-analytics.readonly",
            "https://www.googleapis.com/auth/drive",
        ],
    )


# ---------- Trends ----------
def get_trending_keywords(seed_keywords, target_geos=("US", "GB")):
    all_trending = []
    seeds = [
        k.strip()
        for k in seed_keywords.split(",")
        if k.strip()
    ][:5]

    for geo in target_geos:
        try:
            pytrends = TrendReq(hl="en-US", tz=360)
            pytrends.build_payload(
                seeds,
                timeframe="now 7-d",
                geo=geo.strip(),
            )

            related = pytrends.related_queries()

            for kw in seeds:
                data = related.get(kw, {})
                rising = data.get("rising")

                if rising is not None:
                    all_trending.extend(
                        rising["query"].tolist()[:5]
                    )

        except Exception as e:
            print(f"Trend research ({geo}) fail hui: {e}")

    return list(set(all_trending)) if all_trending else seeds


# ---------- SEO Generation ----------
def generate_seo_content_for_video(
    trending_keywords,
    niche,
    current_title,
    used_titles,
    target_audience="US and UK",
    is_short=False,
):
    avoid_note = (
        "\nYe titles already is run mein use ho chuke hain, "
        f"inse alag banao: {', '.join(used_titles)}"
        if used_titles
        else ""
    )

    shorts_note = (
        "\nYe ek YouTube SHORT hai "
        "— hashtags mein #Shorts zaroor shamil karo."
        if is_short
        else ""
    )

    prompt = f"""Trending keywords: {', '.join(trending_keywords)}
Niche: {niche}
Target audience: {target_audience} viewers
Is video ka current title: "{current_title}"
{avoid_note}{shorts_note}

Isi video (topic wahi rahega, sirf SEO behtar karo) ke liye naya
SEO-optimized title, description, tags, aur hashtags generate karo.
{target_audience} audience ki language style, phrasing, aur cultural
context ke hisab se likho.

Sirf valid JSON return karo:
{{"title": "...", "description": "...", "tags": ["...", "..."], "hashtags": ["#...", "#..."]}}

Title 60 characters se kam, unique, description 2-3 paragraphs,
10-15 relevant tags, 3-5 hashtags."""

    response = requests.post(
        "https://api.groq.com/openai/v1/chat/completions",
        headers={
            "Authorization": f"Bearer {GROQ_API_KEY}",
            "Content-Type": "application/json",
        },
        json={
            "model": "openai/gpt-oss-20b",
            "messages": [
                {"role": "user", "content": prompt}
            ],
            "temperature": 0.8,
            "reasoning_effort": "low",
            "max_tokens": 1500,
        },
        timeout=30,
    )

    if response.status_code >= 400:
        print(
            f"GROQ DEBUG: status={response.status_code} "
            f"body={response.text[:500]}"
        )

    response.raise_for_status()

    raw_text = response.json()["choices"][0]["message"]["content"]
    clean_text = (
        raw_text.replace("```json", "")
        .replace("```", "")
        .strip()
    )

    return json.loads(clean_text)


# ---------- Video Details ----------
def get_recent_videos(youtube, channel_id, max_results=5):
    search_response = youtube.search().list(
        channelId=channel_id,
        part="id",
        order="date",
        maxResults=max_results,
        type="video",
    ).execute()

    return [
        item["id"]["videoId"]
        for item in search_response.get("items", [])
    ]


def get_video_details(youtube, video_id):
    response = youtube.videos().list(
        part="snippet,contentDetails",
        id=video_id,
    ).execute()

    if not response.get("items"):
        return None

    item = response["items"][0]

    return {
        "current_title": item["snippet"]["title"],
        "published_at": item["snippet"]["publishedAt"],
        "duration_seconds": parse_iso8601_duration(
            item["contentDetails"]["duration"]
        ),
    }


def parse_iso8601_duration(duration_str):
    match = re.match(
        r"PT(?:(\d+)H)?(?:(\d+)M)?(?:(\d+)S)?",
        duration_str,
    )

    if not match:
        return 0

    hours, minutes, seconds = (
        int(g) if g else 0
        for g in match.groups()
    )

    return hours * 3600 + minutes * 60 + seconds


# ---------- Analytics: UPDATED ----------
def get_video_analytics(youtube_analytics, channel_id, video_id):
    result = {
        "views": "N/A",
        "impressions": "N/A",
        "ctr": "N/A",
        "avg_view_duration": "N/A",
    }

    try:
        response = youtube_analytics.reports().query(
            ids=f"channel=={channel_id}",
            startDate="2020-01-01",
            endDate=datetime.now(timezone.utc).strftime("%Y-%m-%d"),
            metrics="views,averageViewDuration",
            filters=f"video=={video_id}",
        ).execute()

        rows = response.get("rows", [])

        if not rows:
            print(
                f"Video {video_id}: analytics rows "
                "abhi available nahi."
            )
            return result

        headers = [
            column["name"]
            for column in response.get("columnHeaders", [])
        ]

        values = dict(zip(headers, rows[0]))

        result["views"] = values.get("views", "N/A")
        result["avg_view_duration"] = values.get(
            "averageViewDuration",
            "N/A",
        )

        return result

    except Exception as e:
        print(
            f"Analytics fetch fail hui video "
            f"{video_id} ke liye: {e}"
        )
        return result


# ---------- Tracking Helpers ----------
def get_previous_row(video_id, tracking_file):
    if not os.path.exists(tracking_file):
        return None

    wb = load_workbook(tracking_file)
    ws = wb.active
    previous = None

    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[1] == video_id:
            previous = row

    wb.close()
    return previous


def calculate_delta(current_val, previous_val):
    if not isinstance(current_val, (int, float)):
        return "N/A"

    if not isinstance(previous_val, (int, float)):
        return "N/A"

    diff = current_val - previous_val
    sign = "+" if diff >= 0 else ""

    if isinstance(diff, float):
        return f"{sign}{diff:.2f}"

    return f"{sign}{diff}"


# ---------- Performance Diagnosis ----------
def diagnose_video_performance(
    video_id,
    analytics,
    duration_seconds,
    published_at,
    previous_diagnosis,
):
    try:
        published_dt = datetime.fromisoformat(
            published_at.replace("Z", "+00:00")
        )

        age_days = (
            datetime.now(timezone.utc) - published_dt
        ).days

    except Exception:
        age_days = 0

    if age_days < DIAGNOSIS_MIN_AGE_DAYS:
        return None

    ctr = analytics.get("ctr")
    impressions = analytics.get("impressions")
    avg_duration = analytics.get("avg_view_duration")

    if (
        not isinstance(ctr, (int, float))
        or not isinstance(impressions, (int, float))
    ):
        return None

    category = None
    message = None

    if impressions < LOW_IMPRESSIONS_THRESHOLD:
        category = "REACH"
        message = (
            f"REACH/SEO ISSUE: Video {video_id} ke "
            f"{impressions} impressions hain "
            f"({age_days} din mein). "
            "Keywords/tags ya upload consistency dekhein."
        )

    elif ctr < LOW_CTR_THRESHOLD:
        category = "THUMBNAIL"
        message = (
            f"THUMBNAIL/TITLE ISSUE: Video {video_id} ke "
            f"{impressions} impressions hain, "
            f"CTR sirf {ctr}%. "
            "Thumbnail/title weak lag raha hai."
        )

    elif (
        isinstance(avg_duration, (int, float))
        and duration_seconds > 0
    ):
        retention_percent = (
            avg_duration / duration_seconds
        ) * 100

        if retention_percent < LOW_RETENTION_PERCENT:
            category = "CONTENT"
            message = (
                f"CONTENT/HOOK ISSUE: Video {video_id} pe "
                f"log sirf ~{retention_percent:.0f}% "
                "video dekh rahe hain. "
                "Script/pacing ka masla ho sakta hai."
            )

    if not message:
        return (
            f"Video {video_id}: numbers theek lag rahe hain, "
            "koi major red flag nahi mila."
        )

    if (
        previous_diagnosis
        and category
        and category in previous_diagnosis
    ):
        message += (
            " ESCALATION: Pichli update ke baad bhi "
            "yehi masla hai — sirf text-level change "
            "(title/tags) kaafi nahi. "
            "Thumbnail image quality ya content bhi dekhein."
        )

    return f"[{category}] {message}"


# ---------- Description and Metadata ----------
def build_description_with_hashtags(seo_data):
    description = seo_data.get("description", "")
    hashtags = seo_data.get("hashtags", [])

    if hashtags:
        description = (
            description.rstrip()
            + "\n\n"
            + " ".join(hashtags[:5])
        )

    return description


def update_video_seo(youtube, video_id, seo_data):
    video_response = youtube.videos().list(
        part="snippet",
        id=video_id,
    ).execute()

    if not video_response.get("items"):
        print(f"Video {video_id} nahi mila, skip.")
        return False

    snippet = video_response["items"][0]["snippet"]

    snippet["title"] = seo_data["title"][:100]
    snippet["description"] = build_description_with_hashtags(
        seo_data
    )
    snippet["tags"] = seo_data.get("tags", [])[:15]

    youtube.videos().update(
        part="snippet",
        body={
            "id": video_id,
            "snippet": snippet,
        },
    ).execute()

    print(f"Video {video_id} ka SEO update ho gaya.")
    return True


# ---------- Tracking Sheet ----------
def log_to_tracking_sheet(
    video_id,
    seo_data,
    analytics_before,
    previous_row,
    tracking_file,
    status="Success",
    diagnosis="",
):
    headers = [
        "Date",
        "Video ID",
        "Title Used",
        "Tags Used",
        "Views",
        "Views Delta",
        "Impressions",
        "Impressions Delta",
        "CTR %",
        "CTR Delta",
        "Avg View Duration sec",
        "Status",
        "Diagnosis",
    ]

    if os.path.exists(tracking_file):
        wb = load_workbook(tracking_file)
        ws = wb.active

    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "SEO Tracking"
        ws.append(headers)

    views_delta = "N/A (pehli update)"
    impressions_delta = "N/A (pehli update)"
    ctr_delta = "N/A (pehli update)"

    if previous_row:
        views_delta = calculate_delta(
            analytics_before["views"],
            previous_row[4],
        )
        impressions_delta = calculate_delta(
            analytics_before["impressions"],
            previous_row[6],
        )
        ctr_delta = calculate_delta(
            analytics_before["ctr"],
            previous_row[8],
        )

    ws.append([
        datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"),
        video_id,
        seo_data.get("title", "N/A"),
        ", ".join(seo_data.get("tags", [])),
        analytics_before["views"],
        views_delta,
        analytics_before["impressions"],
        impressions_delta,
        analytics_before["ctr"],
        ctr_delta,
        analytics_before["avg_view_duration"],
        status,
        diagnosis,
    ])

    wb.save(tracking_file)
    wb.close()

    print(f"Tracking sheet mein row add ho gayi: {tracking_file}")


# ---------- Per-Channel Processing ----------
def process_channel(ctx):
    run_summary = [
        f"[{ctx['name']}] Run — "
        f"{datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}"
    ]

    diagnosis_notes = []

    quota = qf.QuotaTracker(
        safety_limit=ctx["quota_safety_limit"]
    )

    if not GROQ_API_KEY:
        run_summary.append(
            "GROQ_API_KEY missing hai — channel skip."
        )
        return run_summary

    print(
        f"[{ctx['name']}] Trending keywords "
        "research ho rahi hai..."
    )

    trending = get_trending_keywords(
        ctx["niche_keywords"],
        target_geos=ctx["target_audience_geos"],
    )

    winning_keywords = gf.get_winning_keywords(
        winning_file=ctx["winning_keywords_file"]
    )

    if winning_keywords:
        trending = list(set(winning_keywords + trending))
        run_summary.append(
            "Winning keywords bhi include kiye: "
            + ", ".join(winning_keywords)
        )

    run_summary.append(
        f"Trending keywords: {', '.join(trending[:8])}"
    )

    surges = gf.check_trending_surge(
        ctx["niche_keywords"],
        ctx["target_audience_geos"],
    )

    if surges:
        run_summary.append("\n--- Surge Alerts ---")
        run_summary.extend(surges)

    if (
        not ctx["youtube_api_key"]
        or not ctx["youtube_channel_id"]
    ):
        run_summary.append(
            "YouTube credentials abhi nahi diye gaye "
            "— sirf research test hua."
        )
        return run_summary

    youtube_read = build(
        "youtube",
        "v3",
        developerKey=ctx["youtube_api_key"],
    )

    video_ids = get_recent_videos(
        youtube_read,
        ctx["youtube_channel_id"],
    )

    quota.add("videos_list")

    run_summary.append(
        f"Videos process karne ke liye mile: {len(video_ids)}"
    )

    if quota.can_afford("search"):
        competitor_titles = gf.track_competitor_titles(
            youtube_read,
            ctx["niche_keywords"],
        )

        quota.add("search")

        if competitor_titles:
            run_summary.append(
                "\n--- Competitor Trending Titles "
                "(pichle 14 din, top-viewed) ---"
            )
            run_summary.extend(
                [f"- {t}" for t in competitor_titles[:5]]
            )

    else:
        run_summary.append(
            "\nQuota limit ke paas "
            "— competitor tracking skip ki."
        )

    if not (
        ctx["oauth_client_id"]
        and ctx["oauth_client_secret"]
        and ctx["oauth_refresh_token"]
    ):
        run_summary.append(
            "OAuth credentials missing hain "
            "— update/analytics/diagnosis skip."
        )
        return run_summary

    creds = get_oauth_credentials(ctx)

    youtube_write = build(
        "youtube",
        "v3",
        credentials=creds,
    )

    youtube_analytics = build(
        "youtubeAnalytics",
        "v2",
        credentials=creds,
    )

    used_titles = []

    # ---------- Calendar-Based Upload ----------
    today_weekday = datetime.now(timezone.utc).weekday()

    shorts_days = [
        int(d)
        for d in ctx["shorts_upload_days"].split(",")
        if d.strip() != ""
    ]

    long_days = [
        int(d)
        for d in ctx["long_upload_days"].split(",")
        if d.strip() != ""
    ]

    upload_jobs = []

    if (
        today_weekday in shorts_days
        and ctx["drive_shorts_toupload"]
        and ctx["drive_shorts_uploaded"]
    ):
        upload_jobs.append((
            "Short",
            ctx["drive_shorts_toupload"],
            ctx["drive_shorts_uploaded"],
            True,
        ))

    if (
        today_weekday in long_days
        and ctx["drive_long_toupload"]
        and ctx["drive_long_uploaded"]
    ):
        upload_jobs.append((
            "Long video",
            ctx["drive_long_toupload"],
            ctx["drive_long_uploaded"],
            False,
        ))

    if upload_jobs:
        drive_service = build(
            "drive",
            "v3",
            credentials=creds,
        )

        for label, to_folder, uploaded_folder, is_short in upload_jobs:
            new_videos = uf.list_new_drive_videos(
                drive_service,
                to_folder,
            )

            run_summary.append(
                f"\n[{label}] Aaj ka calendar din match hua. "
                f"Nayi videos mili: {len(new_videos)}"
            )

            if new_videos:
                video_file = new_videos[0]
                local_path = None

                try:
                    local_path = uf.download_drive_file(
                        drive_service,
                        video_file["id"],
                        video_file["name"],
                    )

                    if is_short:
                        short_info = uf.validate_short_eligibility(local_path)
                        run_summary.append(
                            f"[Short] Eligibility OK: {short_info['duration_seconds']:.1f}s, "
                            f"{short_info['width']}x{short_info['height']}"
                        )

                    seo_data = generate_seo_content_for_video(
                        trending,
                        ctx["niche_keywords"],
                        video_file["name"],
                        used_titles,
                        ctx["target_audience_label"],
                        is_short,
                    )

                    used_titles.append(seo_data["title"])

                    final_description = (
                        build_description_with_hashtags(seo_data)
                    )

                    publish_at = uf.calculate_next_publish_time(
                        ctx["publish_hour_utc"]
                    )

                    new_video_id = uf.upload_video_to_youtube(
                        youtube_write,
                        local_path,
                        seo_data["title"],
                        final_description,
                        seo_data.get("tags", []),
                        publish_at=publish_at,
                        category_id=ctx["category_id"],
                        video_language=ctx["video_language"],
                        made_for_kids=ctx["made_for_kids"],
                        contains_synthetic_media=ctx["contains_synthetic_media"],
                    )

                    uf.move_drive_file(
                        drive_service,
                        video_file["id"],
                        to_folder,
                        uploaded_folder,
                    )

                    uf.cleanup_downloaded_file(local_path)

                    run_summary.append(
                        f"[{label}] Uploaded "
                        f"(scheduled public at {publish_at}): "
                        f"'{video_file['name']}' -> {new_video_id}"
                    )

                except Exception as e:
                    if local_path:
                        uf.cleanup_downloaded_file(local_path)
                    run_summary.append(
                        f"[{label}] Upload fail hui "
                        f"'{video_file['name']}' ke liye: {e}"
                    )

                    print(
                        f"Upload error for "
                        f"{video_file['name']}: {e}"
                    )

    else:
        run_summary.append(
            f"\nAaj ({today_weekday}) calendar ke hisab se "
            "koi upload scheduled nahi hai."
        )

    # ---------- Metadata Refresh ----------
    success_count = 0
    fail_count = 0

    day_of_year = (
        datetime.now(timezone.utc).timetuple().tm_yday
    )

    run_metadata_refresh = (
        ctx["enable_metadata_refresh"]
        and day_of_year % 3 == 0
    )

    if not run_metadata_refresh:
        run_summary.append(
            "\nAaj metadata-refresh wala din nahi hai "
            "(channel config mein disabled hai ya aaj 3-day cycle nahi) "
            "— sirf upload calendar check hua."
        )

    for vid in (video_ids if run_metadata_refresh else []):
        try:
            details = get_video_details(youtube_write, vid)

            if not details:
                continue

            is_short_video = details["duration_seconds"] <= 180

            seo_data = generate_seo_content_for_video(
                trending,
                ctx["niche_keywords"],
                details["current_title"],
                used_titles,
                ctx["target_audience_label"],
                is_short_video,
            )

            used_titles.append(seo_data["title"])

            analytics_before = get_video_analytics(
                youtube_analytics,
                ctx["youtube_channel_id"],
                vid,
            )

            previous_row = get_previous_row(
                vid,
                ctx["tracking_file"],
            )

            previous_diagnosis = (
                previous_row[12]
                if previous_row and len(previous_row) > 12
                else ""
            )

            diagnosis = ""

            diagnosis_result = diagnose_video_performance(
                vid,
                analytics_before,
                details["duration_seconds"],
                details["published_at"],
                previous_diagnosis,
            )

            if diagnosis_result:
                diagnosis = diagnosis_result
                diagnosis_notes.append(diagnosis_result)

            # CTR unavailable ho to learning skip karo.
            ctr = analytics_before.get("ctr")

            if isinstance(ctr, (int, float)):
                gf.update_keyword_learning(
                    vid,
                    seo_data.get("tags", []),
                    ctr,
                    winning_file=ctx["winning_keywords_file"],
                )

            else:
                print(
                    f"Video {vid}: CTR unavailable "
                    "— keyword learning skip."
                )

            if (
                ctx["enable_auto_comment"]
                and quota.can_afford("comment_insert")
            ):
                gf.generate_and_post_comment(
                    youtube_write,
                    vid,
                    ctx["niche_keywords"],
                )

                quota.add("comment_insert")

            if (
                diagnosis
                and "[THUMBNAIL]" in diagnosis
                and "ESCALATION" in diagnosis
            ):
                run_summary.append(
                    f"ACTION NEEDED: Video {vid} pe "
                    "thumbnail issue REPEAT ho raha hai."
                )

            if (
                ctx["enable_caption_improvement"]
                and quota.can_afford("captions_list")
            ):
                try:
                    caption_id = qf.get_auto_caption_track(
                        youtube_write,
                        vid,
                    )

                    quota.add("captions_list")

                    if (
                        caption_id
                        and quota.can_afford("captions_download")
                        and quota.can_afford("captions_insert")
                    ):
                        raw_caption = qf.download_caption_track(
                            youtube_write,
                            caption_id,
                        )

                        quota.add("captions_download")

                        improved_caption = qf.improve_caption_text(
                            raw_caption
                        )

                        qf.upload_improved_caption(
                            youtube_write,
                            vid,
                            improved_caption,
                        )

                        quota.add("captions_insert")

                        run_summary.append(
                            "Caption improve karke upload ho gaya "
                            f"video {vid} ke liye."
                        )

                except Exception as e:
                    print(
                        "Caption improvement fail hui "
                        f"video {vid} ke liye: {e}"
                    )

            updated = update_video_seo(
                youtube_write,
                vid,
                seo_data,
            )

            quota.add("videos_update")

            log_to_tracking_sheet(
                vid,
                seo_data,
                analytics_before,
                previous_row,
                ctx["tracking_file"],
                "Success" if updated else "Skipped",
                diagnosis,
            )

            if updated:
                success_count += 1
            else:
                fail_count += 1

        except Exception as e:
            print(
                f"Video {vid} process karte waqt error: {e}"
            )
            fail_count += 1

    if run_metadata_refresh:
        run_summary.append(
            f"\nMetadata updates successful: {success_count}, "
            f"Failed/Skipped: {fail_count}"
        )

    run_summary.append(quota.status_message())

    if diagnosis_notes:
        run_summary.append(
            "\n--- Diagnosis (7+ din purani videos) ---"
        )
        run_summary.extend(diagnosis_notes)

    # ---------- Community Ideas ----------
    community_ideas = gf.generate_community_post_ideas(
        ctx["niche_keywords"],
        trending,
    )

    if community_ideas:
        run_summary.append(
            "\n--- Community Post Ideas "
            "(verify karke manual post karein) ---"
        )
        run_summary.extend(
            [f"- {idea}" for idea in community_ideas]
        )

    if datetime.now(timezone.utc).weekday() == 0:
        run_summary.append(
            "\n"
            + gf.generate_weekly_summary(
                tracking_file=ctx["tracking_file"]
            )
        )

    return run_summary


# ---------- Main ----------
def main():
    channels = load_channels()

    if not channels:
        print("channels_config.json mein koi channel nahi mila.")
        return

    all_messages = []

    for channel_cfg in channels:
        ctx = build_channel_context(channel_cfg)

        try:
            summary = process_channel(ctx)
            all_messages.append("\n".join(summary))

        except Exception as e:
            error_text = (
                f"[{ctx['name']}] System mein error aayi:\n"
                f"{str(e)}\n\n"
                f"{traceback.format_exc()[:1200]}"
            )

            print(error_text)
            all_messages.append(error_text)

    final_report = "\n\n=====\n\n".join(all_messages)

    print("\n" + "=" * 20 + " RUN SUMMARY " + "=" * 20)
    print(final_report)

    send_whatsapp_message(final_report)


if __name__ == "__main__":
    main()
